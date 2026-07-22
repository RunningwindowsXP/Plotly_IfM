import openpyxl
import sqlite3
import os
from openpyxl.utils import get_column_letter

# ---------------------------------------------
# Pfade
# ---------------------------------------------
excel_path = input("Pfad zur Excel-Datei: ").strip('"')
jahr = input("Auf welches Jahr beziehen sich die Daten? ")

# Spaltenverschiebung
offset = int(input("Spaltenverschiebung eingeben (Basisspalte: H; 0 = keine, 1 = eine nach rechts, -1 = eine nach links): "))

if not os.path.exists(excel_path):
    raise FileNotFoundError(excel_path)

# ---------------------------------------------
# Excel öffnen
# ---------------------------------------------
wb = openpyxl.load_workbook(excel_path, data_only=True)

# ---------------------------------------------
# DB verbinden
# ---------------------------------------------
conn = sqlite3.connect("daten_neu.db")
cur = conn.cursor()

cur.execute(
    """CREATE TABLE IF NOT EXISTS Wirtschaftszweige (
        ID INTEGER PRIMARY KEY,
        Jahr INTEGER,
        Wirtschaftszweig TEXT,
        Umsatzklasse INTEGER,
        Beschäftigtenklasse INTEGER,
        Anzahl INTEGER,
        [Abhängig Beschäftigte] INTEGER,
        [SV-Beschäftigte] INTEGER,
        [Geringfügig Beschäftigte] INTEGER,
        [Umsatz in 1000€] INTEGER
    )"""
)

conn.commit()

cur.execute(
    """SELECT MAX(ID)
    FROM Wirtschaftszweige"""
)

result = cur.fetchone()[0]
id_nr = 1 if result is None else result + 1


# ---------------------------------------------
# Excel einlesen
# ---------------------------------------------
for ws in wb.worksheets[2:]:

    wirtschaftszweig = ws.title

    # Startspalten der 5 Blöcke
    start_spalten = [
        8,   # H
        13,  # M
        18,  # R
        23,  # W
        28   # AB
    ]

    for umsatzklasse, start_spalte in enumerate(start_spalten, start=1):

        block = []

        # Offset berücksichtigen
        start_spalte += offset

        for row in range(9, 14):  # Zeilen 9-13
            zeile = []

            for col in range(start_spalte, start_spalte + 5):
                wert = ws.cell(row=row, column=col).value

                if wert == ".":
                    wert = None

                zeile.append(wert)

            block.append(zeile)


        # Jede Zeile des Blocks wird ein Datensatz
        for beschaeftigtenklasse in range(5):

            werte = block[beschaeftigtenklasse]

            cur.execute("""
                INSERT INTO Wirtschaftszweige
                (
                    ID,
                    Jahr,
                    Wirtschaftszweig,
                    Umsatzklasse,
                    Beschäftigtenklasse,
                    Anzahl,
                    [Abhängig Beschäftigte],
                    [SV-Beschäftigte],
                    [Geringfügig Beschäftigte],
                    [Umsatz in 1000€]
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                id_nr,
                jahr,
                wirtschaftszweig,
                umsatzklasse,
                beschaeftigtenklasse + 1,

                werte[0],
                werte[1],
                werte[2],
                werte[3],
                werte[4],
            ))

            id_nr += 1


conn.commit()
conn.close()

print(f"Fertig! {id_nr - 1} Datensätze importiert.")